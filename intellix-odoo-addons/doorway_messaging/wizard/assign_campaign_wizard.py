# -*- coding: utf-8 -*-
"""Assistant : assigner plusieurs contacts à une campagne WhatsApp / SMS.

Réponse directe au besoin « assigner une campagne avec plusieurs contacts sur
WhatsApp et SMS ». Pré-rempli depuis une sélection multiple de leads CRM ou de
contacts, alimenté manuellement, ou par import d'un fichier CSV / XLSX.
"""
import base64
import csv
import io
import logging
import re

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# Mots-clés d'en-tête reconnus (normalisés sans accents/espaces) -> rôle logique
_HEADER_NAME = {
    "name", "nom", "prenom", "prenoms", "contact", "fullname", "nomcomplet",
    "nomprenom", "client", "destinataire",
}
_HEADER_PHONE = {
    "phone", "telephone", "tel", "mobile", "gsm", "number", "numero",
    "portable", "whatsapp", "msisdn", "cell", "cellulaire", "phonenumber",
    "numerotelephone",
}
_HEADER_EMAIL = {"email", "emails", "mail", "courriel", "adressemail", "eemail"}

_CSV_TEMPLATE = (
    "nom,telephone,email\n"
    "Jean Dupont,+33612345678,jean.dupont@example.com\n"
    "Marie Martin,+33698765432,marie.martin@example.com\n"
)


def _strip_accents(value):
    import unicodedata

    return "".join(
        c
        for c in unicodedata.normalize("NFKD", value or "")
        if not unicodedata.combining(c)
    )


def _norm_header(value):
    base = _strip_accents(str(value or "")).lower()
    return re.sub(r"[^a-z0-9]", "", base)


def _norm_phone(value):
    raw = str(value or "").strip()
    # openpyxl peut renvoyer un float (ex: 33612345678.0)
    if raw.endswith(".0"):
        raw = raw[:-2]
    digits = re.sub(r"[^\d+]", "", raw)
    if not digits:
        return ""
    if digits.startswith("00"):
        digits = "+" + digits[2:]
    # on ne garde qu'un éventuel '+' en tête
    plus = digits.startswith("+")
    digits = re.sub(r"[^\d]", "", digits)
    return ("+" + digits) if plus else digits


class DoorwayAssignCampaignWizard(models.TransientModel):
    _name = "doorway.assign.campaign.wizard"
    _description = "Assigner des contacts à une campagne WhatsApp / SMS"

    mode = fields.Selection(
        [
            ("new", "Nouvelle campagne"),
            ("existing", "Ajouter à une campagne existante"),
        ],
        default="new",
        required=True,
    )
    campaign_id = fields.Many2one(
        "doorway.message.campaign",
        string="Campagne existante",
        domain="[('statut', 'in', ('brouillon', 'planifie'))]",
    )
    name = fields.Char("Nom de la campagne")

    lead_ids = fields.Many2many(
        "crm.lead",
        "doorway_assign_wizard_lead_rel",
        "wizard_id",
        "lead_id",
        string="Leads CRM",
    )
    partner_ids = fields.Many2many(
        "res.partner",
        "doorway_assign_wizard_partner_rel",
        "wizard_id",
        "partner_id",
        string="Contacts",
    )
    phones_manuels = fields.Text("Téléphones supplémentaires (un par ligne)")

    # --- Import de contacts (CSV / XLSX) ---
    import_file = fields.Binary("Fichier de contacts (CSV / XLSX)")
    import_filename = fields.Char("Nom du fichier")
    import_summary = fields.Text("Résultat de l'import", readonly=True)
    csv_template = fields.Binary(
        "Modèle CSV", compute="_compute_csv_template", readonly=True
    )
    csv_template_name = fields.Char(
        "Nom du modèle", compute="_compute_csv_template", readonly=True
    )

    canal_whatsapp = fields.Boolean("WhatsApp", default=True)
    canal_sms = fields.Boolean("SMS")

    whatsapp_template_id = fields.Many2one(
        "doorway.message.template",
        string="Template WhatsApp approuvé",
        domain="[('canal', '=', 'whatsapp'), ('is_twilio_ready', '=', True)]",
    )
    sms_template_id = fields.Many2one(
        "doorway.message.template",
        string="Template SMS",
        domain="[('canal', '=', 'sms')]",
    )
    corps_whatsapp = fields.Text("Message WhatsApp (si pas de template)")
    corps_sms = fields.Text("Message SMS (si pas de template)")

    send_now = fields.Boolean("Envoyer immédiatement", default=False)

    nb_destinataires = fields.Integer(
        "Destinataires avec téléphone", compute="_compute_stats"
    )
    nb_contacts = fields.Integer("Contacts sélectionnés", compute="_compute_stats")
    warning = fields.Text("Avertissement", compute="_compute_warning")

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        active_model = self.env.context.get("active_model")
        active_ids = self.env.context.get("active_ids") or []
        if active_model == "crm.lead" and active_ids:
            res["lead_ids"] = [(6, 0, active_ids)]
        elif active_model == "res.partner" and active_ids:
            res["partner_ids"] = [(6, 0, active_ids)]
        if self.env.context.get("default_canal_sms"):
            res["canal_sms"] = True
        if self.env.context.get("default_canal_whatsapp") is False:
            res["canal_whatsapp"] = False
        return res

    def _iter_phones(self):
        self.ensure_one()
        phones = set()
        for lead in self.lead_ids:
            phone = lead.phone or lead.mobile
            if phone:
                phones.add(re.sub(r"[^\d+]", "", phone))
        for partner in self.partner_ids:
            phone = partner.phone or partner.mobile
            if phone:
                phones.add(re.sub(r"[^\d+]", "", phone))
        for line in (self.phones_manuels or "").splitlines():
            line = re.sub(r"[^\d+]", "", line.strip())
            if line:
                phones.add(line)
        phones.discard("")
        return phones

    @api.depends("lead_ids", "partner_ids", "phones_manuels")
    def _compute_stats(self):
        for wiz in self:
            wiz.nb_contacts = len(wiz.lead_ids) + len(wiz.partner_ids)
            wiz.nb_destinataires = len(wiz._iter_phones())

    def _compute_csv_template(self):
        content = base64.b64encode(_CSV_TEMPLATE.encode("utf-8"))
        for wiz in self:
            wiz.csv_template = content
            wiz.csv_template_name = "modele_import_contacts.csv"

    # ------------------------------------------------------------------
    # Import de contacts
    # ------------------------------------------------------------------
    def _decode_rows(self):
        """Renvoie une liste de dicts {name, phone, email} depuis le fichier.

        Supporte CSV (séparateur ',' ou ';', encodage UTF-8 / latin-1) et XLSX.
        Détecte les colonnes par en-tête souple ; sinon, devine par contenu.
        """
        self.ensure_one()
        if not self.import_file:
            raise UserError(_("Sélectionnez d'abord un fichier à importer."))
        data = base64.b64decode(self.import_file)
        fname = (self.import_filename or "").lower()

        if fname.endswith((".xlsx", ".xlsm", ".xls")):
            table = self._read_xlsx(data)
        else:
            table = self._read_csv(data)

        if not table:
            return []

        # Détection de l'en-tête
        header = table[0]
        roles = self._map_header(header)
        if roles:
            body = table[1:]
        else:
            roles = None
            body = table

        rows = []
        for raw in body:
            if not any((c or "").strip() for c in raw):
                continue
            if roles:
                rec = self._row_by_roles(raw, roles)
            else:
                rec = self._row_by_guess(raw)
            if rec["phone"] or rec["email"] or rec["name"]:
                rows.append(rec)
        return rows

    @staticmethod
    def _read_csv(data):
        text = None
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = data.decode("utf-8", errors="ignore")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        return [list(r) for r in reader]

    @staticmethod
    def _read_xlsx(data):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(
                _("L'import XLSX nécessite la librairie openpyxl. Utilisez un CSV.")
            )
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        table = []
        for row in ws.iter_rows(values_only=True):
            table.append(["" if c is None else str(c) for c in row])
        wb.close()
        return table

    @staticmethod
    def _map_header(header):
        roles = {}
        for idx, cell in enumerate(header):
            key = _norm_header(cell)
            if not key:
                continue
            if key in _HEADER_NAME and "name" not in roles:
                roles["name"] = idx
            elif key in _HEADER_PHONE and "phone" not in roles:
                roles["phone"] = idx
            elif key in _HEADER_EMAIL and "email" not in roles:
                roles["email"] = idx
        # En-tête valable s'il contient au moins téléphone ou email
        if "phone" in roles or "email" in roles:
            return roles
        return {}

    @staticmethod
    def _row_by_roles(raw, roles):
        def cell(role):
            idx = roles.get(role)
            if idx is None or idx >= len(raw):
                return ""
            return (raw[idx] or "").strip()

        return {
            "name": cell("name"),
            "phone": _norm_phone(cell("phone")),
            "email": cell("email").lower(),
        }

    @staticmethod
    def _row_by_guess(raw):
        name, phone, email = "", "", ""
        for cell in raw:
            val = (cell or "").strip()
            if not val:
                continue
            if "@" in val and not email:
                email = val.lower()
            elif re.fullmatch(r"[+\d\s().\-]{6,}", val) and not phone:
                phone = _norm_phone(val)
            elif not name:
                name = val
        return {"name": name, "phone": phone, "email": email}

    def _find_existing(self, phone, email):
        """Cherche un res.partner ou crm.lead existant par téléphone / email."""
        Partner = self.env["res.partner"].sudo()
        Lead = self.env["crm.lead"].sudo()
        sig = re.sub(r"[^\d]", "", phone or "")[-9:]
        if email:
            partner = Partner.search([("email", "=ilike", email)], limit=1)
            if partner:
                return ("partner", partner)
            lead = Lead.search([("email_from", "=ilike", email)], limit=1)
            if lead:
                return ("lead", lead)
        if sig:
            partner = Partner.search(
                ["|", ("phone", "ilike", sig), ("mobile", "ilike", sig)], limit=1
            )
            if partner:
                return ("partner", partner)
            lead = Lead.search(
                ["|", ("phone", "ilike", sig), ("mobile", "ilike", sig)], limit=1
            )
            if lead:
                return ("lead", lead)
        return (None, None)

    def action_import_file(self):
        """Parse le fichier, déduplique et alimente les destinataires."""
        self.ensure_one()
        rows = self._decode_rows()

        existing_partner_ids = set(self.partner_ids.ids)
        existing_lead_ids = set(self.lead_ids.ids)
        existing_phones = {
            re.sub(r"[^\d]", "", p)[-9:]
            for p in self._manual_phone_lines()
            if re.sub(r"[^\d]", "", p)
        }
        new_manual = []
        seen_sig, seen_email = set(), set()

        n_created = n_matched = n_no_phone = n_dup = 0
        Partner = self.env["res.partner"].sudo()

        for rec in rows:
            phone, email, name = rec["phone"], rec["email"], rec["name"]
            sig = re.sub(r"[^\d]", "", phone or "")[-9:]

            if not phone:
                # WhatsApp / SMS exigent un téléphone
                n_no_phone += 1
                continue

            # Déduplication intra-fichier
            if sig and sig in seen_sig:
                n_dup += 1
                continue
            if email and email in seen_email:
                n_dup += 1
                continue
            if sig:
                seen_sig.add(sig)
            if email:
                seen_email.add(email)

            kind, record = self._find_existing(phone, email)
            if kind == "partner":
                if record.id not in existing_partner_ids:
                    existing_partner_ids.add(record.id)
                    n_matched += 1
                else:
                    n_dup += 1
            elif kind == "lead":
                if record.id not in existing_lead_ids:
                    existing_lead_ids.add(record.id)
                    n_matched += 1
                else:
                    n_dup += 1
            elif name or email:
                partner = Partner.create(
                    {
                        "name": name or email or phone,
                        "phone": phone,
                        "email": email or False,
                        "comment": _("Importé via Campagne multi-contacts"),
                    }
                )
                existing_partner_ids.add(partner.id)
                n_created += 1
            else:
                # téléphone seul, sans nom ni email -> liste manuelle
                if sig and sig in existing_phones:
                    n_dup += 1
                else:
                    existing_phones.add(sig)
                    new_manual.append(phone)
                    n_created += 1

        vals = {
            "partner_ids": [(6, 0, list(existing_partner_ids))],
            "lead_ids": [(6, 0, list(existing_lead_ids))],
        }
        if new_manual:
            base = (self.phones_manuels or "").rstrip()
            joined = "\n".join(new_manual)
            vals["phones_manuels"] = (base + "\n" + joined).strip() if base else joined

        summary = _(
            "%(created)s importés, %(matched)s déjà existants, "
            "%(nophone)s sans téléphone ignorés, %(dup)s doublons ignorés."
        ) % {
            "created": n_created,
            "matched": n_matched,
            "nophone": n_no_phone,
            "dup": n_dup,
        }
        vals["import_summary"] = summary
        # le fichier est consommé : on l'efface pour éviter un double import
        vals["import_file"] = False
        vals["import_filename"] = False
        self.write(vals)

        return {
            "type": "ir.actions.act_window",
            "name": _("Campagne multi-contacts (WhatsApp / SMS)"),
            "res_model": "doorway.assign.campaign.wizard",
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    @api.depends(
        "canal_whatsapp", "canal_sms", "whatsapp_template_id", "send_now"
    )
    def _compute_warning(self):
        for wiz in self:
            msgs = []
            if wiz.canal_whatsapp and not wiz.whatsapp_template_id:
                msgs.append(
                    _(
                        "WhatsApp à l'initiative de l'entreprise nécessite un "
                        "template approuvé. Sans template, l'envoi n'aboutira "
                        "que dans la fenêtre de session de 24h."
                    )
                )
            if wiz.send_now:
                msgs.append(
                    _(
                        "« Envoyer immédiatement » déclenchera l'envoi réel via "
                        "Twilio dès validation."
                    )
                )
            wiz.warning = "\n".join(msgs)

    def _manual_phone_lines(self):
        self.ensure_one()
        return [
            line.strip()
            for line in (self.phones_manuels or "").splitlines()
            if line.strip()
        ]

    def action_assign(self):
        self.ensure_one()
        if not (self.canal_whatsapp or self.canal_sms):
            raise UserError(_("Sélectionnez au moins un canal (WhatsApp ou SMS)."))
        if not (self.lead_ids or self.partner_ids or self._manual_phone_lines()):
            raise UserError(_("Sélectionnez au moins un contact ou un téléphone."))

        Campaign = self.env["doorway.message.campaign"]

        if self.mode == "existing":
            if not self.campaign_id:
                raise UserError(_("Choisissez une campagne existante."))
            campaign = self.campaign_id
            vals = {
                "lead_ids": [(4, lead.id) for lead in self.lead_ids],
                "partner_ids": [(4, p.id) for p in self.partner_ids],
            }
            if self.canal_whatsapp:
                vals["canal_whatsapp"] = True
            if self.canal_sms:
                vals["canal_sms"] = True
            if self.whatsapp_template_id:
                vals["whatsapp_template_id"] = self.whatsapp_template_id.id
            if self.sms_template_id:
                vals["sms_template_id"] = self.sms_template_id.id
            if self.corps_whatsapp:
                vals["corps_whatsapp"] = self.corps_whatsapp
            if self.corps_sms:
                vals["corps_sms"] = self.corps_sms
            existing_phones = campaign.phones_manuels or ""
            extra_phones = "\n".join(self._manual_phone_lines())
            if extra_phones:
                vals["phones_manuels"] = (
                    (existing_phones + "\n" + extra_phones).strip()
                )
            campaign.write(vals)
        else:
            if not self.name:
                raise UserError(_("Indiquez un nom de campagne."))
            campaign = Campaign.create(
                {
                    "name": self.name,
                    "canal_whatsapp": self.canal_whatsapp,
                    "canal_sms": self.canal_sms,
                    "whatsapp_template_id": self.whatsapp_template_id.id,
                    "sms_template_id": self.sms_template_id.id,
                    "corps_whatsapp": self.corps_whatsapp,
                    "corps_sms": self.corps_sms,
                    "lead_ids": [(6, 0, self.lead_ids.ids)],
                    "partner_ids": [(6, 0, self.partner_ids.ids)],
                    "phones_manuels": "\n".join(self._manual_phone_lines()),
                    "envoi_immediat": self.send_now,
                    "statut": "brouillon",
                }
            )

        if self.send_now:
            campaign.action_envoyer()

        return {
            "type": "ir.actions.act_window",
            "name": _("Campagne"),
            "res_model": "doorway.message.campaign",
            "res_id": campaign.id,
            "view_mode": "form",
            "target": "current",
        }
